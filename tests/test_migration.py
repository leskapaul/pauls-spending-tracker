"""
Unit tests for spreadsheet migration rules (migrate_history.py):
Rule 1: Non-monthly sheets (e.g. 'Total', '2026') are rollups and do not yield transaction records.
Rule 2: Monthly sheets only parse the first four columns (Date, Description, Amount, Notes).
Rule 3: Bottom/unsectioned transactions default to 'Off-Budget / Uncategorized' unless explicitly
        referenced by a rollup sheet formula.
"""

import os
import tempfile
import openpyxl
import pytest
from migrate_history import (
    is_monthly_sheet,
    categorize_rollup_label,
    extract_rollup_associations,
    parse_monthly_sheet,
    migrate,
)


class TestRule1RollupSheets:
    """Rule 1: Non-monthly sheets provide no raw transactions (only rollups/verifications)."""

    def test_is_monthly_sheet(self):
        # Valid 6-digit monthly sheet names
        assert is_monthly_sheet("202601") is True
        assert is_monthly_sheet("202202") is True
        assert is_monthly_sheet("202512") is True

        # Non-monthly rollup sheets
        assert is_monthly_sheet("Total") is False
        assert is_monthly_sheet("2026") is False
        assert is_monthly_sheet("2025") is False
        assert is_monthly_sheet("Summary") is False
        assert is_monthly_sheet("Notes") is False
        assert is_monthly_sheet("") is False
        assert is_monthly_sheet(None) is False

    def test_rollup_sheets_excluded_from_transactions(self):
        wb = openpyxl.Workbook()
        ws_default = wb.active
        ws_default.title = "Total"
        # Add summary rows to Total sheet
        ws_default.cell(row=1, column=1, value="Months So Far")
        ws_default.cell(row=1, column=2, value=12)
        ws_default.cell(row=2, column=1, value="Groceries Average")
        ws_default.cell(row=2, column=2, value="=DIVIDE(SUM('202601'!C4),1)")

        # Add yearly rollup sheet
        ws_yearly = wb.create_sheet("2026")
        ws_yearly.cell(row=1, column=1, value="Repairs Total")
        ws_yearly.cell(row=1, column=2, value="=SUM('202601'!C7)")

        # Add actual monthly sheet
        ws_m = wb.create_sheet("202601")
        # Row 1: Header
        ws_m.cell(row=1, column=1, value="Date")
        ws_m.cell(row=1, column=2, value="Description")
        ws_m.cell(row=1, column=3, value="Amount")
        ws_m.cell(row=1, column=4, value="Notes")
        # Row 2: Transaction
        ws_m.cell(row=2, column=1, value="2026-01-05")
        ws_m.cell(row=2, column=2, value="Trader Joe's")
        ws_m.cell(row=2, column=3, value=-45.20)
        # Row 3: Transaction
        ws_m.cell(row=3, column=1, value="2026-01-10")
        ws_m.cell(row=3, column=2, value="Whole Foods")
        ws_m.cell(row=3, column=3, value=-80.00)
        # Row 4: Groceries subtotal
        ws_m.cell(row=4, column=2, value="groceries total")
        ws_m.cell(row=4, column=3, value="=SUM(C2:C3)")

        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = os.path.join(tmpdir, "Spending_Test.xlsx")
            wb.save(excel_path)

            df = migrate(excel_path=excel_path, data_dir=tmpdir)

            # Assert all extracted records originate strictly from monthly sheets
            assert len(df) == 2
            assert set(df["sheet_source"].unique()) == {"202601"}
            assert "Total" not in df["sheet_source"].values
            assert "2026" not in df["sheet_source"].values


class TestRule2ColumnScope:
    """Rule 2: Migration considers only the first four columns (Date, Description, Amount, Notes)."""

    def test_strictly_first_four_columns_parsed(self):
        wb = openpyxl.Workbook()
        ws_m = wb.active
        ws_m.title = "202601"

        # Row 2: Categorized Groceries item with columns 5+ populated
        ws_m.cell(row=2, column=1, value="2026-01-15")
        ws_m.cell(row=2, column=2, value="Shoprite Fair Lawn")
        ws_m.cell(row=2, column=3, value=-120.50)
        ws_m.cell(row=2, column=4, value="Weekly groceries")
        # Columns 5+ contain extra reimbursement and split data that must be ignored
        ws_m.cell(row=2, column=5, value="due jing")
        ws_m.cell(row=2, column=6, value=-60.25)
        ws_m.cell(row=2, column=7, value="paid jing via zelle")
        ws_m.cell(row=2, column=8, value="split 50/50")

        # Row 3: Subtotal row
        ws_m.cell(row=3, column=2, value="groceries total")
        ws_m.cell(row=3, column=3, value="=SUM(C2:C2)")

        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = os.path.join(tmpdir, "Spending_Test.xlsx")
            wb.save(excel_path)

            df = migrate(excel_path=excel_path, data_dir=tmpdir)
            assert len(df) == 1
            row = df.iloc[0]

            assert row["date"] == "2026-01-15"
            assert row["description"] == "Shoprite Fair Lawn"
            assert row["amount"] == -120.50
            assert row["notes"] == "Weekly groceries"
            # Ensure columns 5+ were completely ignored and not appended into notes
            assert "due jing" not in row["notes"]
            assert "paid jing" not in row["notes"]
            assert "split" not in row["notes"]

    def test_reimbursement_rows_ignored_and_do_not_override_categories(self):
        wb = openpyxl.Workbook()
        ws_m = wb.active
        ws_m.title = "202601"

        # Row 2: Groceries item
        ws_m.cell(row=2, column=1, value="2026-01-10")
        ws_m.cell(row=2, column=2, value="Trader Joe's")
        ws_m.cell(row=2, column=3, value=-50.00)
        # Row 3: Groceries total
        ws_m.cell(row=3, column=2, value="groceries total")
        ws_m.cell(row=3, column=3, value="=SUM(C2:C2)")

        # Row 4: Discretionary item
        ws_m.cell(row=4, column=1, value="2026-01-12")
        ws_m.cell(row=4, column=2, value="Movie tickets")
        ws_m.cell(row=4, column=3, value=-30.00)
        # Row 5: Discretionary total
        ws_m.cell(row=5, column=2, value="discretionary total")
        ws_m.cell(row=5, column=3, value="=SUM(C4:C4)")

        # Row 6: 'due jing' reimbursement formula referencing C2 and C4
        ws_m.cell(row=6, column=2, value="due jing")
        ws_m.cell(row=6, column=3, value="=SUM(C2,C4)")

        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = os.path.join(tmpdir, "Spending_Test.xlsx")
            wb.save(excel_path)

            df = migrate(excel_path=excel_path, data_dir=tmpdir)
            assert len(df) == 2

            tj_row = df[df["description"] == "Trader Joe's"].iloc[0]
            assert tj_row["category"] == "Groceries"

            movie_row = df[df["description"] == "Movie tickets"].iloc[0]
            assert movie_row["category"] == "Discretionary"

            # 'due jing' must neither be extracted as a transaction nor appear as a category
            assert "due jing" not in df["description"].str.lower().values
            assert "Due Jing" not in df["category"].values


class TestRule3BottomRowCategorization:
    """Rule 3: Unsectioned bottom rows default to Off-Budget unless referenced by a rollup sheet."""

    @pytest.fixture
    def test_workbook_path(self):
        wb = openpyxl.Workbook()

        # 1. Rollup sheet "2026"
        ws_2026 = wb.active
        ws_2026.title = "2026"
        # B10 references C8 as Repairs Total
        ws_2026.cell(row=10, column=1, value="Repairs Total")
        ws_2026.cell(row=10, column=2, value="=SUM('202601'!C8)")
        # B14 references C9 as Bespoke Investment Repairs
        ws_2026.cell(row=14, column=1, value="Bespoke Investment Repairs")
        ws_2026.cell(row=14, column=2, value="=SUM('202601'!C9)")

        # 2. Monthly sheet "202601"
        ws_m = wb.create_sheet("202601")
        # Regular section: Groceries (rows 2-3, subtotal at row 4)
        ws_m.cell(row=2, column=1, value="2026-01-02")
        ws_m.cell(row=2, column=2, value="Shoprite")
        ws_m.cell(row=2, column=3, value=-50.00)

        ws_m.cell(row=3, column=1, value="2026-01-03")
        ws_m.cell(row=3, column=2, value="Trader Joe's")
        ws_m.cell(row=3, column=3, value=-75.00)

        ws_m.cell(row=4, column=2, value="groceries total")
        ws_m.cell(row=4, column=3, value="=SUM(C2:C3)")

        # Bottom row 7: NO category section and NO rollup reference
        # Even with keywords like "attic upgrade", must NOT be presumptuously categorized
        ws_m.cell(row=7, column=1, value="2026-01-20")
        ws_m.cell(row=7, column=2, value="attic upgrade incomplete")
        ws_m.cell(row=7, column=3, value=-1400.00)

        # Bottom row 8: NO section in monthly sheet, BUT referenced in 2026!B10 (Repairs Total)
        ws_m.cell(row=8, column=1, value="2026-01-22")
        ws_m.cell(row=8, column=2, value="rav4 prime side view mirror repair")
        ws_m.cell(row=8, column=3, value=-135.00)

        # Bottom row 9: NO section in monthly sheet, BUT referenced in 2026!B14 (Bespoke Investment Repairs)
        ws_m.cell(row=9, column=1, value="2026-01-25")
        ws_m.cell(row=9, column=2, value="furnace repair for 9049 rental")
        ws_m.cell(row=9, column=3, value=-480.00)

        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = os.path.join(tmpdir, "Spending_Test.xlsx")
            wb.save(excel_path)
            yield excel_path, tmpdir

    def test_unreferenced_bottom_row_is_off_budget(self, test_workbook_path):
        excel_path, tmpdir = test_workbook_path
        df = migrate(excel_path=excel_path, data_dir=tmpdir)

        # Row 7 had no total row and no rollup reference
        unref_row = df[df["description"] == "attic upgrade incomplete"].iloc[0]
        assert unref_row["category"] == "Off-Budget / Uncategorized"

    def test_referenced_bottom_rows_receive_rollup_category(self, test_workbook_path):
        excel_path, tmpdir = test_workbook_path
        df = migrate(excel_path=excel_path, data_dir=tmpdir)

        # Row 8 was referenced by 2026!B10 (Repairs Total)
        repairs_row = df[df["description"] == "rav4 prime side view mirror repair"].iloc[0]
        assert repairs_row["category"] == "Repairs"
        assert repairs_row["amount"] == -135.00

        # Row 9 was referenced by 2026!B14 (Bespoke Investment Repairs)
        rental_row = df[df["description"] == "furnace repair for 9049 rental"].iloc[0]
        assert rental_row["category"] == "Property Maintenance"
        assert rental_row["amount"] == -480.00


class TestRollupLabelMapping:
    """Tests label parsing from rollup sheets into canonical categories."""

    def test_categorize_rollup_label(self):
        assert categorize_rollup_label("Groceries Average") == "Groceries"
        assert categorize_rollup_label("Discretionary Average") == "Discretionary"
        assert categorize_rollup_label("Data Average") == "Data"
        assert categorize_rollup_label("Utilities Average") == "Utilities"
        assert categorize_rollup_label("Repairs Total") == "Repairs"
        assert categorize_rollup_label("Vacations Total") == "Vacations"
        assert categorize_rollup_label("Home Upgrades Total") == "Home Upgrades"
        assert categorize_rollup_label("Bespoke Investment Repairs") == "Property Maintenance"
        assert categorize_rollup_label("Investment Property Upkeep") == "Property Maintenance"
        assert categorize_rollup_label("9049 Property Maintenance") == "Property Maintenance"
        assert categorize_rollup_label("Deficit vs Plan") is None
        assert categorize_rollup_label("") is None
        assert categorize_rollup_label(None) is None
